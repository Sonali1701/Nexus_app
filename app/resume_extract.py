"""Extract candidate contact fields (name, email, phone) from resume files.

Text is extracted locally (PyMuPDF for PDF, python-docx for DOCX). Fields are
found heuristically first — regex for email/phone, top-of-resume scan for the
name. If heuristics leave gaps (or the PDF is a scanned image with no text
layer), an optional Claude call fills them in, controlled by CLAUDE_EXTRACT
in .env: "auto" (default, only when heuristics are incomplete), "always", or
"never". Claude uses whatever Anthropic credentials the environment provides.
"""
import base64
import io
import json
import logging
import re

import fitz  # PyMuPDF
from docx import Document

from . import config

try:
    import anthropic
except ImportError:
    anthropic = None

log = logging.getLogger("resume_extract")

FIELDS = ("firstName", "lastName", "email", "phone")

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?<!\d)(?:\+?1[\s.\-]?)?\(?\d{3}\)?[\s.\-]?\d{3}[\s.\-]?\d{4}(?!\d)")

# Credential/degree suffixes common on (healthcare) resumes: "Jane Doe, RN-BC, CCRN".
# Matched on letters only, so "RN-BC" and "R.N." both normalize into this set.
CREDENTIALS = {"rn", "bsn", "lpn", "lvn", "cna", "np", "aprn", "msn", "dnp",
               "md", "do", "pa", "pac", "pt", "dpt", "ot", "cota", "crna",
               "cma", "rma", "rrt", "cst", "stna", "ccrn", "cen", "cnor",
               "cpn", "rnc", "rnbc", "chpn", "ocn", "pccn", "tcrn", "emt",
               "phd", "mba", "msw", "lcsw", "bls", "acls", "pals", "nrp",
               "tncc", "enpc", "ibclc", "cnm", "whnp", "fnp", "agnp", "pmhnp"}
NAME_PREFIXES = {"dr", "mr", "mrs", "ms", "prof"}
SURNAME_PARTICLES = {"da", "de", "del", "della", "der", "di", "dos", "du",
                     "la", "le", "st", "van", "von"}
STOP_WORDS = {"resume", "curriculum", "vitae", "cv", "summary", "objective",
              "profile", "professional", "contact", "address", "phone", "email",
              "references", "experience", "education", "skills", "nurse",
              "registered", "travel", "licensed", "practitioner", "physician",
              "specialist", "manager", "coordinator", "director", "supervisor",
              "assistant", "technician", "technologist", "therapist",
              "emergency", "department", "critical", "intensive", "pediatric",
              "surgical", "university", "college", "hospital", "health",
              "healthcare", "medical", "center", "centre", "clinic", "staffing",
              "agency", "services", "solutions", "group", "institute", "school",
              "academy", "rehabilitation", "rehab", "senior", "living",
              "united", "states", "inc", "llc", "corp", "company"}

# contact headers are often "Name | City, ST | phone | email" on one line
SEG_SPLIT = re.compile(r"[|•·∙‣◦/]+|\t+| {3,}")
CITY_STATE_RE = re.compile(r",\s*([A-Z]{2})\.?\s*$")
LABEL_RE = re.compile(r"^(?:name|candidate(?:\s+name)?)\s*[:\-]\s*", re.I)
NAME_WORD_RE = re.compile(r"[^\W\d_]+(?:[-'’][^\W\d_]+)*\.?")
INITIAL_RE = re.compile(r"[A-Za-z]\.?")
PHONE_LABEL_RE = re.compile(r"\b(?:cell|mobile|phone|tel|call|contact)\b", re.I)
FAX_RE = re.compile(r"\bfax\b", re.I)


# --- text extraction -------------------------------------------------------

def extract_text(filename: str, content: bytes) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            with fitz.open(stream=content, filetype="pdf") as doc:
                return "\n".join(page.get_text() for page in doc[:3])
        if name.endswith(".docx"):
            doc = Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            return "\n".join(parts)
        if name.endswith(".rtf"):
            txt = content.decode("utf-8", errors="ignore")
            txt = re.sub(r"\\[a-z]+-?\d* ?", " ", txt)
            return txt.replace("{", " ").replace("}", " ")
        # .txt, .doc (old binary .doc gives garbage; heuristics may still find email)
        return content.decode("utf-8", errors="ignore")
    except Exception as e:
        log.warning("text extraction failed for %s: %s", filename, e)
        return ""


# --- heuristics ------------------------------------------------------------

def _format_phone(raw: str) -> str:
    digits = re.sub(r"\D", "", raw)
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    return raw.strip()


def _cred_key(word: str) -> str:
    """Normalize a word for credential/stop-word comparison: letters only."""
    return re.sub(r"[^a-z]", "", word.lower())


def _first_and_last(words):
    """Drop middle names while retaining common multi-word surnames."""
    last_start = len(words) - 1
    while last_start > 1 and _cred_key(words[last_start - 1]) in SURNAME_PARTICLES:
        last_start -= 1
    return words[0], " ".join(words[last_start:])


def _segment_name(seg: str):
    """Return (first, last) if this text segment looks like a person's name."""
    seg = LABEL_RE.sub("", seg.strip())
    if not seg or len(seg) > 45:
        return None
    if EMAIL_RE.search(seg) or any(ch.isdigit() for ch in seg):
        return None
    m = CITY_STATE_RE.search(seg)
    if m and _cred_key(m.group(1)) not in CREDENTIALS:
        return None  # "Chicago, IL" — a city/state line, not a name
    words = [w.strip(",") for w in seg.split() if w.strip(",")]
    while words and _cred_key(words[0]) in NAME_PREFIXES:
        words.pop(0)
    words = [w for w in words if _cred_key(w) not in CREDENTIALS]
    if len(words) > 2:
        # drop middle initials: "John A. Smith" -> John Smith
        trimmed = [w for w in words if not INITIAL_RE.fullmatch(w)]
        if len(trimmed) >= 2:
            words = trimmed
    if not 2 <= len(words) <= 4:
        return None
    for w in words:
        if not NAME_WORD_RE.fullmatch(w):
            return None
        if _cred_key(w) in STOP_WORDS:
            return None
    words = [w.title() if w.isupper() or w.islower() else w for w in words]
    return _first_and_last(words)


def _name_from_lines(text: str):
    for line in text.splitlines()[:15]:
        for seg in SEG_SPLIT.split(line):
            name = _segment_name(seg)
            if name:
                return name
    return None


def _find_phone(text: str):
    """First phone number in the document, preferring lines labeled
    cell/mobile/phone and never taking one off a fax line."""
    first = None
    for line in text.splitlines():
        if FAX_RE.search(line):
            continue
        m = PHONE_RE.search(line)
        if not m:
            continue
        if PHONE_LABEL_RE.search(line):
            return m.group(0)
        if first is None:
            first = m.group(0)
    return first


def _name_from_email(email: str):
    local = email.split("@")[0]
    parts = [p for p in re.split(r"[._\-\d]+", local) if len(p) > 1]
    if len(parts) >= 2:
        return parts[0].title(), " ".join(p.title() for p in parts[1:])
    return None


def _name_from_filename(filename: str):
    base = re.sub(r"\.[^.]+$", "", filename)
    base = re.sub(r"resume|curriculum|vitae|\bcv\b", " ", base, flags=re.I)
    base = re.sub(r"[_\-.()\[\]]+", " ", base)
    base = re.sub(r"\d+", " ", base)
    parts = [p for p in base.split() if p.strip("'").isalpha()]
    if len(parts) >= 2:
        return parts[0].title(), " ".join(p.title() for p in parts[1:])
    return None


def _heuristic(text: str, filename: str) -> dict:
    result = {f: None for f in FIELDS}
    email = EMAIL_RE.search(text)
    if email:
        result["email"] = email.group(0)
    phone = _find_phone(text)
    if phone:
        result["phone"] = _format_phone(phone)
    name = _name_from_lines(text) \
        or (email and _name_from_email(email.group(0))) \
        or _name_from_filename(filename)
    if name:
        result["firstName"], result["lastName"] = name
    return result


# --- Claude fallback -------------------------------------------------------

_claude_client = None
_claude_disabled = False

CONTACT_SCHEMA = {
    "type": "object",
    "properties": {
        "firstName": {"type": ["string", "null"]},
        "lastName": {"type": ["string", "null"]},
        "email": {"type": ["string", "null"]},
        "phone": {"type": ["string", "null"]},
    },
    "required": ["firstName", "lastName", "email", "phone"],
    "additionalProperties": False,
}

PROMPT = (
    "Extract the candidate's own contact details from this resume. "
    "Return null for any field that is not present. Format phone as (123) 456-7890. "
    "Do not use references' or employers' contact details."
)


def claude_available() -> bool:
    return anthropic is not None and not _claude_disabled \
        and config.CLAUDE_EXTRACT != "never" \
        and (_claude_client is not None or _credentials_present())


def _credentials_present() -> bool:
    """Best-effort check so we don't attempt (and log) doomed API calls."""
    import os
    if os.getenv("ANTHROPIC_API_KEY") or os.getenv("ANTHROPIC_AUTH_TOKEN"):
        return True
    # `ant auth login` profile on disk (the SDK resolves these itself)
    cfg = os.getenv("ANTHROPIC_CONFIG_DIR")
    if not cfg:
        cfg = os.path.join(os.getenv("APPDATA", ""), "Anthropic") if os.name == "nt" \
            else os.path.expanduser("~/.config/anthropic")
    return os.path.isdir(os.path.join(cfg, "credentials"))


def _get_claude():
    global _claude_client, _claude_disabled
    if not claude_available():
        return None
    if _claude_client is None:
        if not _credentials_present():
            log.info("No Anthropic credentials found — extraction is heuristics-only")
            _claude_disabled = True
            return None
        try:
            # Zero-arg client: resolves ANTHROPIC_API_KEY / auth token / profile
            _claude_client = anthropic.Anthropic()
        except Exception as e:
            log.info("Claude extraction unavailable: %s", e)
            _claude_disabled = True
            return None
    return _claude_client


def _claude_extract(text: str, pdf_bytes: bytes | None) -> dict | None:
    """Ask Claude for the contact fields; pdf_bytes is sent as a document
    when the PDF has no text layer (scanned image)."""
    global _claude_disabled
    client = _get_claude()
    if client is None:
        return None

    if pdf_bytes is not None:
        content = [
            {"type": "document",
             "source": {"type": "base64", "media_type": "application/pdf",
                        "data": base64.standard_b64encode(pdf_bytes).decode()}},
            {"type": "text", "text": PROMPT},
        ]
    else:
        content = [{"type": "text", "text": f"{PROMPT}\n\n<resume>\n{text[:8000]}\n</resume>"}]

    try:
        response = client.messages.create(
            model=config.ANTHROPIC_MODEL,
            max_tokens=1024,
            output_config={
                "effort": "low",
                "format": {"type": "json_schema", "schema": CONTACT_SCHEMA},
            },
            messages=[{"role": "user", "content": content}],
        )
        if response.stop_reason == "refusal":
            return None
        raw = next((b.text for b in response.content if b.type == "text"), None)
        return json.loads(raw) if raw else None
    except Exception as e:
        if isinstance(e, anthropic.AuthenticationError) or "authentication" in str(e).lower():
            log.info("Anthropic auth failed — disabling Claude extraction for this run")
            _claude_disabled = True
        else:
            log.warning("Claude extraction failed: %s", e)
        return None


# --- public API ------------------------------------------------------------

def extract_fields(filename: str, content: bytes) -> dict:
    """Return {filename, firstName, lastName, email, phone, source, missing}."""
    text = extract_text(filename, content)
    result = _heuristic(text, filename)
    source = "heuristic"

    is_scanned_pdf = filename.lower().endswith(".pdf") and len(text.strip()) < 50
    missing = [f for f in FIELDS if not result[f]]
    want_claude = (
        config.CLAUDE_EXTRACT == "always"
        or (config.CLAUDE_EXTRACT == "auto" and (missing or is_scanned_pdf))
    )

    if want_claude and claude_available():
        claude = _claude_extract(text, content if is_scanned_pdf else None)
        if claude:
            source = "claude" if is_scanned_pdf else "heuristic+claude"
            if config.CLAUDE_EXTRACT == "always" or is_scanned_pdf:
                # Claude read the actual document — prefer its answer
                for f in FIELDS:
                    result[f] = claude.get(f) or result[f]
            else:
                # fill only the gaps; regex hits stay authoritative
                for f in FIELDS:
                    result[f] = result[f] or claude.get(f)

    out = {"filename": filename, "source": source,
           "missing": [f for f in FIELDS if not result[f]]}
    out.update({f: result[f] or "" for f in FIELDS})
    if is_scanned_pdf and not any(result[f] for f in FIELDS):
        out["note"] = "No text layer found (scanned image?) and AI extraction unavailable"
    return out
