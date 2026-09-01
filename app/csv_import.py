"""Parse and safely filter supported candidate CSV files."""
import ast
import csv
import hashlib
import io
import re
import unicodedata
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook

try:
    import xlrd
except ImportError:  # pragma: no cover - dependency is present in packaged builds
    xlrd = None


LICENSE_MIN_LIKELIHOOD = 5
INDEED_MIN_LIKELIHOOD = 3
# Kept for compatibility with callers/tests that refer to the original constant.
MIN_LIKELIHOOD = LICENSE_MIN_LIKELIHOOD
MAX_CSV_BYTES = 25 * 1024 * 1024
MAX_CSV_ROWS = 10_000
MAX_CSV_FILES = 100
MAX_TOTAL_CSV_BYTES = 100 * 1024 * 1024
SUPPORTED_TABLE_EXTENSIONS = {".csv", ".xlsx", ".xls"}

LICENSE_REQUIRED_COLUMNS = {
    "input_name",
    "input_region",
    "input_preserved_License Type",
    "input_preserved_License #",
    "input_preserved_License Issue Date",
    "input_preserved_License Exp Date",
    "status",
    "likelihood",
    "full_name",
    "first_name",
    "last_name",
    "mobile_phone",
    "matched",
}
INDEED_REQUIRED_COLUMNS = {
    "Name",
    "Location",
    "Headline",
    "Email",
    "Phone",
    "PDL likelihood",
    "PDL matched inputs",
    "PDL job title",
    "PDL company",
    "Source URL",
}
# Kept as the original license schema for compatibility with existing callers.
REQUIRED_COLUMNS = LICENSE_REQUIRED_COLUMNS

PERSONAL_EMAIL_DOMAINS = {
    "aol.com",
    "fastmail.com",
    "gmail.com",
    "googlemail.com",
    "hotmail.com",
    "icloud.com",
    "live.com",
    "mail.com",
    "msn.com",
    "outlook.com",
    "proton.me",
    "protonmail.com",
    "yahoo.com",
    "ymail.com",
}

US_STATE_NAMES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR",
    "california": "CA", "colorado": "CO", "connecticut": "CT",
    "delaware": "DE", "district of columbia": "DC", "florida": "FL",
    "georgia": "GA", "hawaii": "HI", "idaho": "ID", "illinois": "IL",
    "indiana": "IN", "iowa": "IA", "kansas": "KS", "kentucky": "KY",
    "louisiana": "LA", "maine": "ME", "maryland": "MD",
    "massachusetts": "MA", "michigan": "MI", "minnesota": "MN",
    "mississippi": "MS", "missouri": "MO", "montana": "MT",
    "nebraska": "NE", "nevada": "NV", "new hampshire": "NH",
    "new jersey": "NJ", "new mexico": "NM", "new york": "NY",
    "north carolina": "NC", "north dakota": "ND", "ohio": "OH",
    "oklahoma": "OK", "oregon": "OR", "pennsylvania": "PA",
    "rhode island": "RI", "south carolina": "SC", "south dakota": "SD",
    "tennessee": "TN", "texas": "TX", "utah": "UT", "vermont": "VT",
    "virginia": "VA", "washington": "WA", "west virginia": "WV",
    "wisconsin": "WI", "wyoming": "WY",
}
US_STATE_CODES = frozenset(US_STATE_NAMES.values())

# Ordered from specific credentials to broader profession terms. Headline is
# always checked first; the PDL job title is only a fallback when the headline
# has no confident match.
PROFESSION_RULES = (
    ("CRNA", re.compile(
        r"\bcrna\b|certified registered nurse anesthetist|nurse anesthetist|\banesthetist\b",
        re.I)),
    ("Nurse Practitioner", re.compile(
        r"nurse practitioner|\bcrnp\b|\bfnp(?:-bc|-c)?\b|\baprn\b|\bnp-c\b",
        re.I)),
    ("Physician Assistant", re.compile(r"physician assistant|\bpa-c\b", re.I)),
    ("LPN/LVN", re.compile(
        r"licensed practical nurse|licensed vocational nurse|\blpn\b|\blvn\b",
        re.I)),
    ("CNA", re.compile(r"certified nursing assistant|\bcna\b", re.I)),
    ("RN", re.compile(r"registered nurse|\brn\b", re.I)),
    ("Patient Care Technician", re.compile(
        r"patient care tech(?:nician)?|\bpct\b", re.I)),
    ("Medical Assistant", re.compile(r"medical assistant|\bcma\b", re.I)),
    ("Respiratory Therapy", re.compile(
        r"respiratory therap(?:ist|y)|\brrt\b", re.I)),
    ("Physical Therapy", re.compile(r"physical therap(?:ist|y)|\bdpt\b", re.I)),
    ("Occupational Therapy", re.compile(
        r"occupational therap(?:ist|y)|\botr\b|\bcota\b", re.I)),
    ("Social Services", re.compile(r"social worker|\blcsw\b|\bmsw\b", re.I)),
    ("Midwife", re.compile(r"certified nurse midwife|\bcnm\b|\bmidwife\b", re.I)),
)

# Canonical candidate fields exposed by the browser's column mapper. Aliases
# are intentionally broad; final row validation still prevents unsafe uploads.
MAPPING_FIELDS = (
    ("fullName", "Full name", "core", (
        "full name", "candidate name", "contact name", "name", "input name")),
    ("firstName", "First name", "core", (
        "first name", "firstname", "first", "given name", "given", "forename")),
    ("middleName", "Middle name", "optional", (
        "middle name", "middlename", "middle initial")),
    ("lastName", "Last name", "core", (
        "last name", "lastname", "last", "surname", "family name")),
    ("personalEmail", "Personal email (preferred)", "core", (
        "personal email", "personal email address", "home email", "private email",
        "recommended personal email", "pdl personal email")),
    ("email", "Email (fallback)", "core", (
        "email", "email address", "e mail", "primary email", "work email",
        "business email", "contact email")),
    ("mobilePhone", "Mobile phone (preferred)", "core", (
        "mobile phone", "mobile phone number", "mobile", "cell phone", "cellphone",
        "cell phone number", "pdl mobile phone")),
    ("phone", "Phone (fallback)", "core", (
        "phone", "phone number", "primary phone", "telephone", "contact number",
        "contact phone", "home phone", "work phone")),
    ("location", "Combined city/state", "core", (
        "location", "candidate location", "city state", "city and state",
        "current location", "address location")),
    ("state", "State", "core", (
        "state", "state code", "region", "province", "input region",
        "license state", "licensed state")),
    ("city", "City", "optional", (
        "city", "locality", "town", "input locality")),
    ("addressLine1", "Street address", "optional", (
        "street address", "address", "address line 1", "address1", "street",
        "input street address")),
    ("zip", "ZIP/postal code", "optional", (
        "zip", "zip code", "zipcode", "postal code", "postcode",
        "input postal code")),
    ("profession", "Profession", "core", (
        "profession", "occupation", "discipline", "profession name",
        "license type", "input preserved license type")),
    ("headline", "Professional headline", "core", (
        "headline", "professional headline", "candidate headline", "indeed headline")),
    ("jobTitle", "Job title", "core", (
        "job title", "current job title", "current title", "pdl job title",
        "position", "role", "title")),
    ("company", "Company", "optional", (
        "company", "company name", "employer", "current company", "pdl company")),
    ("likelihood", "Likelihood/confidence", "optional", (
        "pdl likelihood", "likelihood", "confidence", "confidence score",
        "match confidence", "match score", "score")),
    ("linkedin", "LinkedIn/profile URL", "optional", (
        "linkedin", "linkedin url", "linkedin profile", "profile url")),
    ("sourceUrl", "Source URL", "optional", (
        "source url", "indeed url", "candidate url", "source link")),
    ("licenseType", "License type", "optional", (
        "license type", "credential type", "input preserved license type")),
    ("licenseNumber", "License number", "optional", (
        "license number", "license no", "license #", "credential number",
        "input preserved license #")),
    ("licenseIssueDate", "License issue date", "optional", (
        "license issue date", "issue date", "issued date",
        "input preserved license issue date")),
    ("licenseExpirationDate", "License expiration date", "optional", (
        "license expiration date", "license expiry date", "expiration date",
        "expiry date", "input preserved license exp date")),
)

MAPPING_FIELD_KEYS = {field[0] for field in MAPPING_FIELDS}


class CSVImportError(ValueError):
    pass


def _text(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def _normalized_name(value):
    value = unicodedata.normalize("NFKD", _text(value)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def _list(value):
    try:
        parsed = ast.literal_eval(value) if value else []
        return parsed if isinstance(parsed, list) else []
    except (SyntaxError, ValueError):
        return []


def _structured_emails(row):
    return [
        item for item in _list(row.get("emails"))
        if isinstance(item, dict) and _text(item.get("address"))
    ]


def _best_email(row):
    """Prefer every available personal-email source before a work address."""
    recommended = _text(row.get("recommended_personal_email")).lower()
    if recommended:
        return recommended
    for item in _list(row.get("personal_emails")):
        if isinstance(item, str) and _text(item):
            return _text(item).lower()
    structured = _structured_emails(row)
    for item in structured:
        if _text(item.get("type")).lower() == "personal":
            return _text(item["address"]).lower()
    work = _text(row.get("work_email")).lower()
    if work:
        return work
    if structured:
        return _text(structured[0]["address"]).lower()
    return ""


def _first_value(row, keys):
    for key in keys:
        value = _text(row.get(key))
        if value:
            return value
    return ""


def _email_type(email):
    domain = email.rsplit("@", 1)[-1].lower() if "@" in email else ""
    return "personal" if domain in PERSONAL_EMAIL_DOMAINS else "provided fallback"


def _best_indeed_email(row):
    personal = _first_value(row, (
        "Personal Email", "Personal email", "personal_email",
        "recommended_personal_email", "PDL personal email",
    )).lower()
    if personal:
        return personal, "personal"
    email = _text(row.get("Email")).lower()
    return email, _email_type(email)


def _phone(value):
    digits = re.sub(r"\D", "", _text(value))
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    if 10 <= len(digits) <= 15:
        return "+" + digits
    return ""


def _best_indeed_phone(row):
    mobile = _first_value(row, (
        "Mobile Phone", "Mobile phone", "mobile_phone", "PDL mobile phone",
    ))
    if mobile:
        return _phone(mobile), "mobile"
    return _phone(row.get("Phone")), "provided fallback"


def _name_case(value):
    return _text(value).title()


def _decode(content):
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise CSVImportError("CSV encoding is not supported; save it as UTF-8 CSV")


def _split_name(value):
    raw = _text(value)
    if "," in raw:
        last, given = (part.strip() for part in raw.split(",", 1))
        parts = given.split()
        if last and parts:
            return (_name_case(parts[0]), _name_case(" ".join(parts[1:])),
                    _name_case(last))
    parts = raw.split()
    if len(parts) < 2:
        return "", "", ""
    return _name_case(parts[0]), _name_case(" ".join(parts[1:-1])), _name_case(parts[-1])


def _parse_indeed_likelihood(value):
    raw = _text(value)
    if not raw:
        return None
    match = re.search(r"\d+", raw)
    return int(match.group()) if match else -1


def _parse_us_location(value):
    """Return ``(city, state_code)`` for a U.S. city/state location."""
    location = _text(value)
    if "," not in location:
        return "", ""
    city, state = (part.strip() for part in location.rsplit(",", 1))
    state = re.sub(r"\s+\d{5}(?:-\d{4})?$", "", state).strip()
    normalized_state = re.sub(r"\s+", " ", state.replace(".", "")).lower()
    state_code = (
        state.upper() if state.upper() in US_STATE_CODES
        else US_STATE_NAMES.get(normalized_state, "")
    )
    return (city, state_code) if city and state_code else ("", "")


def _detect_profession(headline, pdl_job_title):
    for source, value in (("headline", headline), ("PDL job title", pdl_job_title)):
        text_value = _text(value)
        if not text_value:
            continue
        for profession_name, pattern in PROFESSION_RULES:
            if pattern.search(text_value):
                return profession_name, source
    return "", ""


def _unique_headers(values):
    """Create non-empty unique labels without discarding duplicate columns."""
    headers = []
    counts = Counter()
    for index, value in enumerate(values, start=1):
        base = _text(value) or f"Column {index}"
        counts[base] += 1
        headers.append(base if counts[base] == 1 else f"{base} [{counts[base]}]")
    return headers


def _table_from_matrix(filename, matrix, sheet_name=""):
    rows = list(matrix)
    header_index = next((
        index for index, row in enumerate(rows)
        if any(_text(value) for value in row)
    ), None)
    if header_index is None:
        raise CSVImportError(f"{filename} is empty")
    raw_headers = list(rows[header_index])
    while raw_headers and not _text(raw_headers[-1]):
        raw_headers.pop()
    if not raw_headers:
        raise CSVImportError(f"{filename} does not contain a header row")
    headers = _unique_headers(raw_headers)
    parsed_rows = []
    for physical_row, values in enumerate(rows[header_index + 1:], start=header_index + 2):
        values = list(values[:len(headers)])
        values.extend([""] * (len(headers) - len(values)))
        if not any(_text(value) for value in values):
            continue
        parsed_rows.append({
            "sourceRow": physical_row,
            "values": dict(zip(headers, values)),
        })
        if len(parsed_rows) > MAX_CSV_ROWS:
            raise CSVImportError(
                f"{filename} has more than the {MAX_CSV_ROWS:,}-row limit")
    if not parsed_rows:
        raise CSVImportError(f"{filename} contains headers but no data rows")
    return {
        "filename": filename,
        "sheetName": sheet_name,
        "headers": headers,
        "rows": parsed_rows,
    }


def _read_csv_table(filename, content):
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:16_384], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        matrix = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error as exc:
        raise CSVImportError(f"Could not read {filename}: {exc}") from exc
    return _table_from_matrix(filename, matrix)


def _read_xlsx_table(filename, content):
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise CSVImportError(f"Could not read Excel file {filename}: {exc}") from exc
    try:
        for worksheet in workbook.worksheets:
            matrix = list(worksheet.iter_rows(values_only=True, max_row=MAX_CSV_ROWS + 2))
            if any(any(_text(value) for value in row) for row in matrix):
                return _table_from_matrix(filename, matrix, worksheet.title)
    finally:
        workbook.close()
    raise CSVImportError(f"Excel file {filename} has no non-empty worksheets")


def _read_xls_table(filename, content):
    if xlrd is None:
        raise CSVImportError(
            "Legacy .xls support is unavailable; install xlrd or save the file as .xlsx")
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        for worksheet in workbook.sheets():
            matrix = [worksheet.row_values(index) for index in range(
                min(worksheet.nrows, MAX_CSV_ROWS + 2))]
            if any(any(_text(value) for value in row) for row in matrix):
                return _table_from_matrix(filename, matrix, worksheet.name)
    except Exception as exc:
        raise CSVImportError(f"Could not read Excel file {filename}: {exc}") from exc
    finally:
        try:
            workbook.release_resources()
        except (NameError, AttributeError):
            pass
    raise CSVImportError(f"Excel file {filename} has no non-empty worksheets")


def read_table(filename, content):
    """Read a CSV/XLSX/XLS upload into a common header-and-row structure."""
    if not content:
        raise CSVImportError(f"{filename} is empty")
    if len(content) > MAX_CSV_BYTES:
        raise CSVImportError(f"{filename} is over the 25 MB import limit")
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_TABLE_EXTENSIONS:
        raise CSVImportError("Choose only .csv, .xlsx, or .xls files")
    if extension == ".csv":
        return _read_csv_table(filename, content)
    if extension == ".xlsx":
        return _read_xlsx_table(filename, content)
    return _read_xls_table(filename, content)


def _normalized_header(value):
    value = re.sub(r"\s+\[\d+\]$", "", _text(value))
    return _normalized_name(value)


def _header_alias_score(header, aliases):
    normalized = _normalized_header(header)
    if not normalized:
        return 0
    header_tokens = set(normalized.split())
    best = 0
    for alias in aliases:
        alias = _normalized_header(alias)
        if normalized == alias:
            best = max(best, 120)
            continue
        alias_tokens = set(alias.split())
        if len(alias_tokens) >= 2 and alias_tokens <= header_tokens:
            best = max(best, 96)
        elif len(header_tokens) >= 2 and header_tokens <= alias_tokens:
            best = max(best, 88)
        similarity = SequenceMatcher(None, normalized, alias).ratio()
        if similarity >= .9:
            best = max(best, 86)
    return best


def _column_examples(table):
    examples = {header: [] for header in table["headers"]}
    for item in table["rows"][:50]:
        for header, value in item["values"].items():
            value = _text(value)
            if value and value not in examples[header] and len(examples[header]) < 3:
                examples[header].append(value[:80])
    return examples


def _value_score(field, values):
    values = [value for value in values if value][:20]
    if not values:
        return 0
    if field == "email":
        valid = sum(bool(re.search(r"[^@\s]+@[^@\s]+\.[^@\s]+", value)) for value in values)
        return 84 if valid / len(values) >= .7 else 0
    if field == "phone":
        valid = sum(bool(_phone(value)) for value in values)
        return 84 if valid / len(values) >= .7 else 0
    if field == "state":
        valid = sum(bool(_state_code(value)) for value in values)
        return 82 if valid / len(values) >= .8 else 0
    return 0


def suggest_mapping(table):
    """Return a one-to-one best-effort canonical-to-source column mapping."""
    examples = _column_examples(table)
    candidates = []
    for field, _label, _tier, aliases in MAPPING_FIELDS:
        for header in table["headers"]:
            score = _header_alias_score(header, aliases)
            # Content inference is reserved for generic fallback fields; a
            # value alone cannot prove that an email is personal or a phone mobile.
            if field in {"email", "phone", "state"}:
                score = max(score, _value_score(field, examples[header]))
            if score >= 80:
                candidates.append((score, field, header))
    mapping = {field: "" for field in MAPPING_FIELD_KEYS}
    used_headers = set()
    for _score, field, header in sorted(candidates, reverse=True):
        if not mapping[field] and header not in used_headers:
            mapping[field] = header
            used_headers.add(header)
    return mapping


def _mapping_issues(headers, mapping):
    issues = []
    headers = set(headers)
    for field, header in mapping.items():
        if field in MAPPING_FIELD_KEYS and header and header not in headers:
            issues.append(f"mapped column {header!r} is not present")
    if not mapping.get("fullName") and not (
            mapping.get("firstName") and mapping.get("lastName")):
        issues.append("map Full name, or both First name and Last name")
    if not (mapping.get("personalEmail") or mapping.get("email")):
        issues.append("map a Personal email or Email column")
    if not (mapping.get("mobilePhone") or mapping.get("phone")):
        issues.append("map a Mobile phone or Phone column")
    if not (mapping.get("location") or mapping.get("state")):
        issues.append("map a Combined city/state or State column")
    if not (mapping.get("profession") or mapping.get("headline")
            or mapping.get("jobTitle")):
        issues.append("map a Profession, Professional headline, or Job title column")
    return issues


def _group_id(headers):
    signature = "\x1f".join(_normalized_header(header) for header in headers)
    return hashlib.sha256(signature.encode("utf-8")).hexdigest()[:12]


def _is_legacy_csv(filename, headers):
    if Path(filename).suffix.lower() != ".csv":
        return ""
    header_set = set(headers)
    if INDEED_REQUIRED_COLUMNS <= header_set:
        return "Indeed candidate matches"
    if LICENSE_REQUIRED_COLUMNS <= header_set:
        return "PDL-enriched license export"
    return ""


def inspect_many(named_contents):
    """Inspect layouts and return grouped automatic column suggestions."""
    files = list(named_contents)
    if not files:
        raise CSVImportError("Choose at least one CSV or Excel file")
    if len(files) > MAX_CSV_FILES:
        raise CSVImportError(f"Choose no more than {MAX_CSV_FILES} files at once")
    if sum(len(content) for _, content in files) > MAX_TOTAL_CSV_BYTES:
        raise CSVImportError("Selected files exceed the 100 MB combined limit")

    groups = {}
    recognized = []
    for index, (filename, content) in enumerate(files):
        table = read_table(filename, content)
        legacy_label = _is_legacy_csv(filename, table["headers"])
        if legacy_label:
            recognized.append({"filename": filename, "schemaLabel": legacy_label})
            continue
        group_id = _group_id(table["headers"])
        group = groups.get(group_id)
        if not group:
            suggested = suggest_mapping(table)
            group = {
                "id": group_id,
                "files": [],
                "columns": table["headers"],
                "examples": _column_examples(table),
                "suggestedMapping": suggested,
                "issues": _mapping_issues(table["headers"], suggested),
                "sheetNames": [],
            }
            groups[group_id] = group
        group["files"].append({"index": index, "filename": filename})
        if table.get("sheetName") and table["sheetName"] not in group["sheetNames"]:
            group["sheetNames"].append(table["sheetName"])

    public_fields = [
        {"key": key, "label": label, "tier": tier}
        for key, label, tier, _aliases in MAPPING_FIELDS
    ]
    group_list = list(groups.values())
    return {
        "fields": public_fields,
        "groups": group_list,
        "recognizedFiles": recognized,
        "canAutoMap": all(not group["issues"] for group in group_list),
    }


def _state_code(value):
    value = _text(value)
    value = re.sub(r"\s+\d{5}(?:-\d{4})?$", "", value).strip()
    normalized = re.sub(r"\s+", " ", value.replace(".", "")).lower()
    if value.upper() in US_STATE_CODES:
        return value.upper()
    if normalized in US_STATE_NAMES:
        return US_STATE_NAMES[normalized]
    words = re.findall(r"[A-Za-z]+", value)
    for word in reversed(words):
        if word.upper() in US_STATE_CODES:
            return word.upper()
    return ""


def _generic_location(location, state, city):
    state_code = _state_code(state)
    city_value = _text(city)
    location_value = _text(location)
    if location_value:
        parsed_city, parsed_state = _parse_us_location(location_value)
        if parsed_state:
            state_code = state_code or parsed_state
            city_value = city_value or parsed_city
        elif not state_code:
            state_code = _state_code(location_value)
        if not city_value and state_code:
            city_value = re.sub(
                rf"(?:,|\s)\s*(?:{re.escape(state_code)}|[A-Za-z ]+)\s*(?:\d{{5}}(?:-\d{{4}})?)?$",
                "", location_value, flags=re.I).strip(" ,")
    return _name_case(city_value), state_code


def _mapped(row, mapping, field):
    header = mapping.get(field)
    return row.get(header) if header else ""


def _first_valid_email(*values):
    for value in values:
        match = re.search(r"[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+", _text(value))
        if match:
            return match.group(0).lower()
    return ""


def _generic_profession(profession, headline, job_title):
    explicit = _text(profession)
    if explicit:
        detected, _source = _detect_profession(explicit, "")
        return (detected or explicit, "profession column")
    return _detect_profession(headline, job_title)


def _parse_generic_table(table, mapping):
    issues = _mapping_issues(table["headers"], mapping)
    if issues:
        raise CSVImportError(
            f"Column mapping is incomplete for {table['filename']}: " + "; ".join(issues))

    safe = []
    rejected = Counter()
    seen_email = set()
    seen_phone = set()
    for item in table["rows"]:
        row = item["values"]
        first_name = _name_case(_mapped(row, mapping, "firstName"))
        middle_name = _name_case(_mapped(row, mapping, "middleName"))
        last_name = _name_case(_mapped(row, mapping, "lastName"))
        if not first_name or not last_name:
            first_name, split_middle, last_name = _split_name(
                _mapped(row, mapping, "fullName"))
            middle_name = middle_name or split_middle
        if not first_name or not last_name:
            rejected["missing parsed first or last name"] += 1
            continue

        personal_raw = _mapped(row, mapping, "personalEmail")
        email = _first_valid_email(personal_raw, _mapped(row, mapping, "email"))
        if not email:
            rejected["missing or invalid email"] += 1
            continue
        email_type = "personal" if _first_valid_email(personal_raw) == email else _email_type(email)

        mobile_raw = _mapped(row, mapping, "mobilePhone")
        mobile_phone = _phone(mobile_raw)
        phone = mobile_phone or _phone(_mapped(row, mapping, "phone"))
        if not phone:
            rejected["missing or invalid phone"] += 1
            continue
        phone_source = "mobile" if mobile_phone else "provided fallback"

        city, state_code = _generic_location(
            _mapped(row, mapping, "location"),
            _mapped(row, mapping, "state"),
            _mapped(row, mapping, "city"),
        )
        if not state_code:
            rejected["missing or invalid U.S. state"] += 1
            continue

        headline = _text(_mapped(row, mapping, "headline"))
        job_title = _text(_mapped(row, mapping, "jobTitle"))
        profession_name, profession_source = _generic_profession(
            _mapped(row, mapping, "profession"), headline, job_title)
        if not profession_name:
            rejected["profession could not be confidently inferred"] += 1
            continue

        likelihood = _parse_indeed_likelihood(_mapped(row, mapping, "likelihood"))
        if likelihood == -1:
            rejected["invalid confidence"] += 1
            continue
        if likelihood is not None and likelihood < INDEED_MIN_LIKELIHOOD:
            rejected[f"confidence below {INDEED_MIN_LIKELIHOOD}"] += 1
            continue

        phone_key = re.sub(r"\D", "", phone)
        if email in seen_email or phone_key in seen_phone:
            rejected["duplicate email or phone inside file"] += 1
            continue
        seen_email.add(email)
        seen_phone.add(phone_key)

        linkedin = _text(_mapped(row, mapping, "linkedin"))
        if linkedin and not linkedin.startswith(("http://", "https://")):
            linkedin = "https://" + linkedin
        safe.append({
            "sourceType": "generic_table",
            "sourceRow": item["sourceRow"],
            "firstName": first_name,
            "middleName": middle_name,
            "lastName": last_name,
            "email": email,
            "emailType": email_type,
            "phone": phone,
            "phoneSource": phone_source,
            "addressLine1": _text(_mapped(row, mapping, "addressLine1")),
            "city": city,
            "stateCode": state_code,
            "zip": _text(_mapped(row, mapping, "zip")),
            "licenseType": _text(_mapped(row, mapping, "licenseType")),
            "licenseNumber": _text(_mapped(row, mapping, "licenseNumber")),
            "licenseIssueDate": _text(_mapped(row, mapping, "licenseIssueDate")),
            "licenseExpirationDate": _text(_mapped(
                row, mapping, "licenseExpirationDate")),
            "linkedin": linkedin,
            "likelihood": likelihood,
            "headline": headline,
            "pdlJobTitle": job_title,
            "pdlCompany": _text(_mapped(row, mapping, "company")),
            "sourceUrl": _text(_mapped(row, mapping, "sourceUrl")),
            "professionName": profession_name,
            "professionSource": profession_source,
        })

    total = len(table["rows"])
    return ({
        "schema": "generic_table",
        "schemaLabel": "custom mapped CSV/Excel",
        "totalRows": total,
        "safeRows": len(safe),
        "rejectedRows": total - len(safe),
        "minimumLikelihood": INDEED_MIN_LIKELIHOOD,
        "blankLikelihoodPolicy": "allowed when confidence is not supplied",
        "safeByState": dict(Counter(row["stateCode"] for row in safe)),
        "safeByProfession": dict(Counter(row["professionName"] for row in safe)),
        "rejectedByReason": dict(rejected),
    }, safe)


def _parse_indeed_rows(reader):
    safe = []
    rejected = Counter()
    seen_email = set()
    seen_phone = set()
    total = 0

    for source_row, row in enumerate(reader, start=2):
        total += 1
        if total > MAX_CSV_ROWS:
            raise CSVImportError(f"CSV has more than the {MAX_CSV_ROWS:,}-row limit")

        likelihood = _parse_indeed_likelihood(row.get("PDL likelihood"))
        if likelihood == -1:
            rejected["invalid PDL confidence"] += 1
            continue
        if likelihood is not None and likelihood < INDEED_MIN_LIKELIHOOD:
            rejected[f"confidence below {INDEED_MIN_LIKELIHOOD}"] += 1
            continue
        matched = {
            part.strip().lower()
            for part in _text(row.get("PDL matched inputs")).split(",")
            if part.strip()
        }
        if likelihood is not None and "name" not in matched:
            rejected["name was not a PDL match signal"] += 1
            continue

        city, state_code = _parse_us_location(row.get("Location"))
        if not city or not state_code:
            rejected["location is not a supported U.S. state"] += 1
            continue
        headline = _text(row.get("Headline"))

        first_name, middle_name, last_name = _split_name(row.get("Name"))
        if not first_name or not last_name:
            rejected["missing parsed first or last name"] += 1
            continue
        email, email_type = _best_indeed_email(row)
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            rejected["missing or invalid email"] += 1
            continue
        phone, phone_source = _best_indeed_phone(row)
        if not phone:
            rejected["missing or invalid phone"] += 1
            continue

        profession_name, profession_source = _detect_profession(
            headline, row.get("PDL job title"))
        if not profession_name:
            rejected["profession could not be confidently inferred"] += 1
            continue

        phone_key = re.sub(r"\D", "", phone)
        if email in seen_email or phone_key in seen_phone:
            rejected["duplicate email or phone inside CSV"] += 1
            continue
        seen_email.add(email)
        seen_phone.add(phone_key)

        safe.append({
            "sourceType": "indeed_matches",
            "sourceRow": source_row,
            "firstName": first_name,
            "middleName": middle_name,
            "lastName": last_name,
            "email": email,
            "emailType": email_type,
            "phone": phone,
            "phoneSource": phone_source,
            "addressLine1": "",
            "city": _name_case(city),
            "stateCode": state_code,
            "zip": "",
            "licenseType": "",
            "licenseNumber": "",
            "licenseIssueDate": "",
            "licenseExpirationDate": "",
            "linkedin": "",
            "likelihood": likelihood,
            "headline": headline,
            "pdlJobTitle": _text(row.get("PDL job title")),
            "pdlCompany": _text(row.get("PDL company")),
            "sourceUrl": _text(row.get("Source URL")),
            "professionName": profession_name,
            "professionSource": profession_source,
        })

    return ({
        "schema": "indeed_matches",
        "schemaLabel": "Indeed candidate matches",
        "totalRows": total,
        "safeRows": len(safe),
        "rejectedRows": total - len(safe),
        "minimumLikelihood": INDEED_MIN_LIKELIHOOD,
        "blankLikelihoodPolicy": (
            "allowed only for complete U.S. contacts"
        ),
        "safeByState": dict(Counter(row["stateCode"] for row in safe)),
        "safeByProfession": dict(Counter(row["professionName"] for row in safe)),
        "rejectedByReason": dict(rejected),
    }, safe)


def parse_and_filter(content):
    """Return ``(summary, safe_records)`` for either supported CSV schema."""
    if not content:
        raise CSVImportError("CSV file is empty")
    if len(content) > MAX_CSV_BYTES:
        raise CSVImportError("CSV file is over the 25 MB import limit")

    reader = csv.DictReader(io.StringIO(_decode(content)))
    headers = set(reader.fieldnames or [])
    if INDEED_REQUIRED_COLUMNS <= headers:
        return _parse_indeed_rows(reader)

    missing = sorted(LICENSE_REQUIRED_COLUMNS - headers)
    if missing:
        raise CSVImportError(
            "CSV format is not supported. For the Kentucky-license format, missing: "
            + ", ".join(missing)
        )

    safe = []
    rejected = Counter()
    seen_email = set()
    seen_phone = set()
    seen_license = set()
    total = 0

    for source_row, row in enumerate(reader, start=2):
        total += 1
        if total > MAX_CSV_ROWS:
            raise CSVImportError(f"CSV has more than the {MAX_CSV_ROWS:,}-row limit")

        if _text(row.get("status")) != "200":
            rejected["no enrichment match"] += 1
            continue
        if _normalized_name(row.get("input_name")) != _normalized_name(row.get("full_name")):
            rejected["enriched name does not exactly match input name"] += 1
            continue
        if "name" not in _list(row.get("matched")):
            rejected["input name was not an enrichment match signal"] += 1
            continue
        try:
            likelihood = int(_text(row.get("likelihood")) or 0)
        except ValueError:
            likelihood = 0
        if likelihood < LICENSE_MIN_LIKELIHOOD:
            rejected[f"confidence below {LICENSE_MIN_LIKELIHOOD}"] += 1
            continue
        if _text(row.get("input_region")).upper() != "KY":
            rejected["not a Kentucky record"] += 1
            continue
        if _text(row.get("input_preserved_License Type")).upper() != "RN":
            rejected["not an RN license"] += 1
            continue

        email = _best_email(row)
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            rejected["missing or invalid email"] += 1
            continue
        phone = _phone(row.get("mobile_phone"))
        if not phone:
            rejected["missing or invalid phone"] += 1
            continue

        license_number = _text(row.get("input_preserved_License #"))
        issue_date = _text(row.get("input_preserved_License Issue Date"))
        expiration_date = _text(row.get("input_preserved_License Exp Date"))
        if not license_number or not issue_date or not expiration_date:
            rejected["incomplete license details"] += 1
            continue
        phone_key = re.sub(r"\D", "", phone)
        if email in seen_email or phone_key in seen_phone or license_number in seen_license:
            rejected["duplicate contact or license inside CSV"] += 1
            continue
        seen_email.add(email)
        seen_phone.add(phone_key)
        seen_license.add(license_number)

        first_name = _name_case(row.get("first_name"))
        last_name = _name_case(row.get("last_name"))
        if not first_name or not last_name:
            rejected["missing parsed first or last name"] += 1
            continue

        linkedin = _text(row.get("linkedin_url"))
        if linkedin and not linkedin.startswith(("http://", "https://")):
            linkedin = "https://" + linkedin

        safe.append({
            "sourceType": "pdl_license",
            "sourceRow": source_row,
            "firstName": first_name,
            "middleName": _name_case(row.get("middle_name")),
            "lastName": last_name,
            "email": email,
            "emailType": _email_type(email),
            "phone": phone,
            "phoneSource": "mobile",
            "addressLine1": _text(row.get("input_street_address")),
            "city": _name_case(row.get("input_locality")),
            "stateCode": _text(row.get("input_region")).upper(),
            "zip": _text(row.get("input_postal_code")),
            "licenseType": _text(row.get("input_preserved_License Type")).upper(),
            "licenseNumber": license_number,
            "licenseIssueDate": issue_date,
            "licenseExpirationDate": expiration_date,
            "linkedin": linkedin,
            "likelihood": likelihood,
        })

    summary = {
        "schema": "pdl_license",
        "schemaLabel": "PDL-enriched Kentucky RN licenses",
        "totalRows": total,
        "safeRows": len(safe),
        "rejectedRows": total - len(safe),
        "minimumLikelihood": LICENSE_MIN_LIKELIHOOD,
        "safeByState": dict(Counter(row["stateCode"] for row in safe)),
        "rejectedByReason": dict(rejected),
    }
    return summary, safe


def parse_and_filter_many(named_contents, mappings=None):
    """Parse mapped CSV/Excel files and deduplicate their accepted rows."""
    files = list(named_contents)
    if not files:
        raise CSVImportError("Choose at least one CSV or Excel file")
    if len(files) > MAX_CSV_FILES:
        raise CSVImportError(f"Choose no more than {MAX_CSV_FILES} files at once")
    if sum(len(content) for _, content in files) > MAX_TOTAL_CSV_BYTES:
        raise CSVImportError("Selected files exceed the 100 MB combined limit")

    mappings = mappings or {}
    summaries = []
    parsed_summaries = []
    records = []
    schemas = set()
    for filename, content in files:
        table = read_table(filename, content)
        legacy_label = _is_legacy_csv(filename, table["headers"])
        if legacy_label:
            summary, file_records = parse_and_filter(content)
        else:
            group_id = _group_id(table["headers"])
            suggested = suggest_mapping(table)
            supplied = mappings.get(group_id)
            if supplied is not None and not isinstance(supplied, dict):
                raise CSVImportError(f"Invalid column mapping for {filename}")
            mapping = dict(suggested)
            if supplied is not None:
                mapping.update({
                    key: _text(value)
                    for key, value in supplied.items()
                    if key in MAPPING_FIELD_KEYS
                })
            summary, file_records = _parse_generic_table(table, mapping)
        schemas.add(summary["schema"])
        parsed_summaries.append(summary)
        summaries.append({
            "filename": filename,
            "sheetName": table.get("sheetName", ""),
            "schemaLabel": summary["schemaLabel"],
            "totalRows": summary["totalRows"],
            "safeRowsBeforeCrossFileDeduplication": summary["safeRows"],
            "rejectedRows": summary["rejectedRows"],
        })
        for record in file_records:
            records.append({**record, "sourceFile": filename})

    # Keep the strongest occurrence when the same email, phone, or license is
    # present in multiple selected files. Personal/mobile contact sources and a
    # higher PDL score win; original file/row order breaks ties.
    ranked = sorted(
        enumerate(records),
        key=lambda item: (
            item[1].get("emailType") == "personal",
            item[1].get("phoneSource") == "mobile",
            item[1].get("likelihood") if item[1].get("likelihood") is not None else -1,
            -item[0],
        ),
        reverse=True,
    )
    kept = []
    seen_email = set()
    seen_phone = set()
    seen_license = set()
    cross_file_duplicates = 0
    for original_index, record in ranked:
        email = record["email"].lower()
        phone = re.sub(r"\D", "", record["phone"])
        license_number = _text(record.get("licenseNumber"))
        if (email in seen_email or phone in seen_phone
                or (license_number and license_number in seen_license)):
            cross_file_duplicates += 1
            continue
        seen_email.add(email)
        seen_phone.add(phone)
        if license_number:
            seen_license.add(license_number)
        kept.append((original_index, record))
    records = [record for _, record in sorted(kept)]

    rejected = Counter()
    for summary in parsed_summaries:
        rejected.update(summary["rejectedByReason"])
    if cross_file_duplicates:
        rejected["duplicate contact across selected CSVs"] += cross_file_duplicates

    schema = next(iter(schemas)) if len(schemas) == 1 else "mixed"
    first_summary = parsed_summaries[0]
    total_rows = sum(item["totalRows"] for item in summaries)
    labels = list(dict.fromkeys(
        summary["schemaLabel"] for summary in parsed_summaries))
    thresholds = {
        summary.get("minimumLikelihood") for summary in parsed_summaries
        if summary.get("minimumLikelihood") is not None
    }
    summary = {
        "schema": schema,
        "schemaLabel": " + ".join(labels),
        "fileCount": len(files),
        "totalRows": total_rows,
        "safeRows": len(records),
        "rejectedRows": total_rows - len(records),
        "minimumLikelihood": min(thresholds) if thresholds else None,
        "likelihoodPolicy": (
            f"confidence {next(iter(thresholds))}+ when supplied"
            if len(thresholds) == 1
            else "format-specific confidence thresholds"
        ),
        "safeByState": dict(Counter(row["stateCode"] for row in records)),
        "safeByProfession": dict(Counter(
            row.get("professionName", "RN") for row in records)),
        "rejectedByReason": dict(rejected),
        "files": summaries,
    }
    blank_policies = list(dict.fromkeys(
        summary.get("blankLikelihoodPolicy") for summary in parsed_summaries
        if summary.get("blankLikelihoodPolicy")))
    if blank_policies:
        summary["blankLikelihoodPolicy"] = "; ".join(blank_policies)
    return summary, records


def candidate_payload(record, defaults):
    """Map one filtered CSV record to the documented Nexus Candidate API."""
    is_license_record = record.get("sourceType") == "pdl_license"
    state_code = _text(record.get("stateCode")).upper()
    state_id = defaults.get("stateIds", {}).get(state_code)
    if state_id is None and state_code == "KY":
        # Backward-compatible fallback for older callers and license tests.
        state_id = defaults.get("stateId")
    if state_id is None:
        raise ValueError(f"No Nexus state mapping is available for {state_code or 'this row'}")

    profession_mapping = defaults.get("professionMappings", {}).get(
        record.get("professionName"), {})
    profession_id = profession_mapping.get("professionId", defaults["professionId"])
    specialty_id = profession_mapping.get("specialtyId", defaults["specialtyId"])

    if is_license_record:
        highlights = (
            f"{state_code} {record['licenseType']} license #{record['licenseNumber']}; "
            f"issued {record['licenseIssueDate']}; expires {record['licenseExpirationDate']}. "
            f"Source: PDL-enriched KY license list "
            f"(confidence {record['likelihood']}/10)."
        )
    elif record.get("sourceType") == "generic_table":
        confidence = (
            f"confidence {record['likelihood']}/10"
            if record.get("likelihood") is not None
            else "no confidence supplied"
        )
        role = record.get("pdlJobTitle") or record.get("headline") \
            or record.get("professionName") or "not supplied"
        company = record.get("pdlCompany")
        employment = f" at {company}" if company else ""
        license_detail = ""
        if record.get("licenseNumber"):
            license_detail = (
                f" License: {record.get('licenseType') or 'type not supplied'} "
                f"#{record['licenseNumber']}; issued "
                f"{record.get('licenseIssueDate') or 'not supplied'}; expires "
                f"{record.get('licenseExpirationDate') or 'not supplied'}."
            )
        highlights = (
            f"Spreadsheet import from {record.get('sourceFile') or 'uploaded file'} "
            f"row {record.get('sourceRow')}. Role: {role}{employment}; {confidence}. "
            f"Detected profession: {record.get('professionName') or 'not supplied'} "
            f"from {record.get('professionSource') or 'mapped column'}; "
            f"Nexus mapping: {profession_mapping.get('professionName', 'batch default')} / "
            f"{profession_mapping.get('specialtyName', 'batch default')}. "
            f"Email: {record.get('emailType', 'provided fallback')}; "
            f"phone: {record.get('phoneSource', 'provided fallback')}."
            f"{license_detail}"
        )
    else:
        confidence = (
            f"PDL confidence {record['likelihood']}/10"
            if record.get("likelihood") is not None
            else "no PDL confidence supplied"
        )
        role = record.get("pdlJobTitle") or record.get("headline") or "not supplied"
        company = record.get("pdlCompany")
        employment = f" at {company}" if company else ""
        highlights = (
            f"Indeed candidate match: {record.get('headline') or 'headline not supplied'}. "
            f"PDL role: {role}{employment}; {confidence}. "
            f"Detected profession: {record.get('professionName') or 'not supplied'} "
            f"from {record.get('professionSource') or 'CSV'}; "
            f"Nexus mapping: {profession_mapping.get('professionName', 'batch default')} / "
            f"{profession_mapping.get('specialtyName', 'batch default')}. "
            f"Email: {record.get('emailType', 'provided fallback')}; "
            f"phone: {record.get('phoneSource', 'provided fallback')}."
        )

    payload = {
        "firstName": record["firstName"],
        "lastName": record["lastName"],
        "primaryEmail": record["email"],
        "phone": record["phone"],
        "cellPhone": record["phone"],
        "statusId": defaults["statusId"],
        "referralSourceId": defaults["referralSourceId"],
        "professionIds": [profession_id],
        "specialtyIds": [specialty_id],
        "primarySpecialtyId": specialty_id,
        "jobTypeIds": defaults["jobTypeIds"],
        "stateId": state_id,
        "countryId": defaults["countryId"],
        "addressLine1": record["addressLine1"] or None,
        "city": record["city"] or None,
        "zip": record["zip"] or None,
        "candidateHighlights": highlights,
        "sendMassEmails": False,
        "sendMassSms": False,
    }
    if is_license_record:
        payload["licensedStateIds"] = [state_id]
    if record.get("middleName"):
        payload["middleName"] = record["middleName"]
    if record.get("linkedin"):
        payload["linkedin"] = record["linkedin"]
    if defaults.get("candidatePRNStatusId"):
        payload["candidatePRNStatusId"] = defaults["candidatePRNStatusId"]
    return payload
