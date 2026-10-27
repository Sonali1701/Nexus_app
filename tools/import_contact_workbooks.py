"""Prepare and resume a specialty-aware import of the requested contact workbooks."""

import argparse
import hashlib
import json
import logging
import os
import re
import sqlite3
import sys
import time
from collections import Counter
from pathlib import Path

from dotenv import dotenv_values
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


WORKBOOK_NAMES = (
    "Leads_1_Nursing_FINAL2.xlsx",
    "Leads_1_Nursing_CLEAN_2.xlsx",
    "Leads_2_Allied_and_Other_CLEAN_2.xlsx",
    "Leads_2_Allied_and_Other_CLEAN_2 - Copy.xlsx",
    "Merged_20_Healthcare_Contacts_Cleaned.xlsx",
    "Merged_10_Healthcare_Contacts_Cleaned.xlsx",
    "Merged_Allied_Healthcare_Contacts_Cleaned.xlsx",
    "Merged_Candidate_Contacts.xlsx",
    "Siddi_Cleaned_Healthcare_Candidates.xlsx",
    "Consolidated_20_Files_US_Healthcare_Contacts.xlsx",
    "Consolidated_20_Healthcare_Contacts_Oct24.xlsx",
    "Consolidated_Healthcare_Contacts_New_Batch.xlsx",
    "Consolidated_New_20_Healthcare_Contacts.xlsx",
    "Consolidated_US_Healthcare_Contacts.xlsx",
    "Healthcare_Contacts_20_Files_Cleaned.xlsx",
)


def load_app(env_path):
    if not env_path.is_file():
        raise FileNotFoundError(env_path)
    os.environ.update({key: value for key, value in dotenv_values(env_path).items()
                       if value is not None})
    from app import csv_import
    from app.nexus_client import client, NexusError
    from app.specialty_mapping import SpecialtyResolver, normalize
    return csv_import, client, NexusError, SpecialtyResolver, normalize


def load_lookup_client(env_path):
    from app import config
    from app.nexus_client import NexusClient

    credentials = dotenv_values(env_path)
    if credentials.get("NEXUS_ORG_CODE") != config.NEXUS_ORG_CODE:
        raise ValueError("Lookup credentials must belong to the upload tenant")
    username = credentials.get("NEXUS_USERNAME")
    password = credentials.get("NEXUS_PASSWORD")
    if not username or not password:
        raise ValueError("Lookup credentials require a Nexus username and password")

    class LookupClient(NexusClient):
        def _fetch_token(self):
            return self._post_token({"grant_type": "password", "username": username,
                                     "password": password,
                                     "organizationCode": config.NEXUS_ORG_CODE})

    return LookupClient()


def open_database(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute("PRAGMA synchronous=NORMAL")
    connection.executescript("""
        CREATE TABLE IF NOT EXISTS contacts (
            id INTEGER PRIMARY KEY,
            record_json TEXT NOT NULL,
            score INTEGER NOT NULL,
            first_key TEXT NOT NULL,
            last_key TEXT NOT NULL,
            state_code TEXT NOT NULL,
            profession_id INTEGER,
            specialty_id INTEGER,
            conflict INTEGER NOT NULL DEFAULT 0,
            merged_into INTEGER,
            status TEXT NOT NULL DEFAULT 'pending',
            candidate_id INTEGER,
            error TEXT
        );
        CREATE TABLE IF NOT EXISTS identifiers (
            kind TEXT NOT NULL,
            value TEXT NOT NULL,
            contact_id INTEGER NOT NULL,
            PRIMARY KEY (kind, value)
        );
        CREATE TABLE IF NOT EXISTS files (
            name TEXT PRIMARY KEY,
            bytes INTEGER NOT NULL,
            modified_ns INTEGER NOT NULL,
            digest TEXT NOT NULL,
            rows INTEGER NOT NULL,
            accepted INTEGER NOT NULL,
            identical_to TEXT
        );
        CREATE INDEX IF NOT EXISTS contacts_queue
            ON contacts (status, merged_into, specialty_id);
    """)
    return connection


def workbook_digest(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def compatible_names(first, second):
    return (first["last_key"] == second["last_key"]
            and first["first_key"][:1] == second["first_key"][:1])


def record_score(record):
    mapping = record.get("nexusMapping") or {}
    source_priority = {"original specialty": 3, "specialty": 2,
                       "profession title": 1}.get(mapping.get("matchSource"), 0)
    return (1000000 if mapping else 0) + source_priority * 100000 \
        + (10000 if record["email"] else 0) \
        + (1000 if record["phoneSource"] == "mobile" else 0) \
        + (100 if record["city"] else 0)


def contact_values(record, csv_import, normalize):
    return {
        "record_json": json.dumps(record, ensure_ascii=False),
        "score": record_score(record),
        "first_key": normalize(record["firstName"]),
        "last_key": normalize(record["lastName"]),
        "state_code": record["stateCode"],
        "profession_id": (record.get("nexusMapping") or {}).get("professionId"),
        "specialty_id": (record.get("nexusMapping") or {}).get("specialtyId"),
    }


def row_to_record(headers, values, source_file, source_row, resolver, csv_import):
    row = dict(zip(headers, values))
    first_name = csv_import._name_case(row.get("First Name"))
    last_name = csv_import._name_case(row.get("Last Name"))
    email = csv_import._first_valid_email(row.get("Email"))
    mobile = csv_import._phone(row.get("Cell Number"))
    phone = mobile or csv_import._phone(row.get("Phone Number"))
    state_code = csv_import._state_code(row.get("State"))
    role = csv_import._text(row.get("Profession/Job title")
                            or row.get("Profession title"))
    if not first_name or not last_name or not state_code or not role \
            or not (email or phone):
        return None
    specialty = csv_import._text(row.get("Specialty"))
    original_specialty = csv_import._text(row.get("Original Specialty"))
    mapping = resolver.resolve(role, specialty, original_specialty)
    return {
        "sourceType": "generic_table",
        "sourceFile": source_file,
        "sourceRow": source_row,
        "firstName": first_name,
        "middleName": "",
        "lastName": last_name,
        "email": email,
        "emailType": csv_import._email_type(email) if email else "",
        "phone": phone,
        "phoneSource": "mobile" if mobile else "provided fallback" if phone else "",
        "addressLine1": "",
        "city": csv_import._name_case(row.get("City")),
        "stateCode": state_code,
        "zip": "",
        "licenseType": "",
        "licenseNumber": "",
        "licenseIssueDate": "",
        "licenseExpirationDate": "",
        "linkedin": "",
        "likelihood": None,
        "headline": "",
        "pdlJobTitle": "",
        "pdlCompany": "",
        "sourceUrl": "",
        "professionName": role,
        "professionSource": "profession column",
        "sourceProfession": role,
        "sourceSpecialty": specialty,
        "originalSpecialty": original_specialty,
        "nexusMapping": mapping,
    }


def identity_keys(record, normalize):
    name = normalize(record["firstName"] + " " + record["lastName"])
    digits = re.sub(r"\D", "", record["phone"])
    keys = []
    if record["email"]:
        keys.append(("email", record["email"].casefold()))
    if digits:
        keys.append(("phone_name", digits + ":" + name))
    return keys


def merge_record(connection, record, csv_import, normalize):
    values = contact_values(record, csv_import, normalize)
    keys = identity_keys(record, normalize)
    matches = sorted({row[0] for kind, value in keys
                      for row in connection.execute(
                          "SELECT contact_id FROM identifiers WHERE kind=? AND value=?",
                          (kind, value))})
    if not matches:
        cursor = connection.execute(
            "INSERT INTO contacts (record_json, score, first_key, last_key, "
            "state_code, profession_id, specialty_id) VALUES (?, ?, ?, ?, ?, ?, ?)",
            tuple(values.values()))
        contact_id = cursor.lastrowid
    else:
        contact_id = matches[0]
        current = connection.execute("SELECT * FROM contacts WHERE id=?",
                                     (contact_id,)).fetchone()
        conflict = bool(current["conflict"]) or not compatible_names(current, values)
        for other_id in matches[1:]:
            other = connection.execute("SELECT * FROM contacts WHERE id=?",
                                       (other_id,)).fetchone()
            conflict = conflict or not compatible_names(current, other) \
                or bool(other["conflict"])
            if other["score"] > current["score"]:
                current = other
            connection.execute("UPDATE identifiers SET contact_id=? WHERE contact_id=?",
                               (contact_id, other_id))
            connection.execute("UPDATE contacts SET merged_into=? WHERE id=?",
                               (contact_id, other_id))
        if values["score"] > current["score"]:
            chosen = values
        else:
            chosen = {key: current[key] for key in values}
        connection.execute(
            "UPDATE contacts SET record_json=:record_json, score=:score, "
            "first_key=:first_key, last_key=:last_key, state_code=:state_code, "
            "profession_id=:profession_id, specialty_id=:specialty_id, "
            "conflict=:conflict WHERE id=:id",
            {**chosen, "conflict": int(conflict or bool(current["conflict"])),
             "id": contact_id})
    for kind, value in keys:
        connection.execute(
            "INSERT OR REPLACE INTO identifiers (kind, value, contact_id) "
            "VALUES (?, ?, ?)", (kind, value, contact_id))


def prepare(connection, downloads, resolver, csv_import, normalize):
    for name in WORKBOOK_NAMES:
        path = downloads / name
        if not path.is_file():
            raise FileNotFoundError(path)
        stat = path.stat()
        existing = connection.execute("SELECT * FROM files WHERE name=?",
                                      (name,)).fetchone()
        if existing:
            if existing["bytes"] != stat.st_size or existing["modified_ns"] != stat.st_mtime_ns:
                raise ValueError(f"{name} changed since staging; use a fresh database")
            print(f"already prepared: {name}", flush=True)
            continue
        digest = workbook_digest(path)
        identical = connection.execute("SELECT name FROM files WHERE digest=? LIMIT 1",
                                       (digest,)).fetchone()
        if identical:
            with connection:
                connection.execute(
                    "INSERT INTO files VALUES (?, ?, ?, ?, 0, 0, ?)",
                    (name, stat.st_size, stat.st_mtime_ns, digest, identical["name"]))
            print(f"identical to {identical['name']}: {name}", flush=True)
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = 0
        accepted = 0
        try:
            with connection:
                for sheet in workbook.worksheets:
                    iterator = sheet.iter_rows(values_only=True)
                    headers = [str(value or "").strip() for value in next(iterator, ())]
                    for source_row, values in enumerate(iterator, start=2):
                        if not any(value is not None for value in values):
                            continue
                        rows += 1
                        record = row_to_record(headers, values, name, source_row,
                                               resolver, csv_import)
                        if record is None:
                            continue
                        accepted += 1
                        merge_record(connection, record, csv_import, normalize)
                connection.execute(
                    "INSERT INTO files VALUES (?, ?, ?, ?, ?, ?, NULL)",
                    (name, stat.st_size, stat.st_mtime_ns, digest, rows, accepted))
        finally:
            workbook.close()
        print(f"prepared: {name}: {rows} rows, {accepted} valid", flush=True)


def show_status(connection):
    files = connection.execute(
        "SELECT COUNT(*), COALESCE(SUM(rows),0), COALESCE(SUM(accepted),0) FROM files"
    ).fetchone()
    counts = Counter({status: count for status, count in connection.execute(
        "SELECT CASE WHEN conflict=1 THEN 'identity_review' "
        "WHEN specialty_id IS NULL THEN 'specialty_review' ELSE status END, "
        "COUNT(*) FROM contacts WHERE merged_into IS NULL GROUP BY 1")})
    print(json.dumps({"files_prepared": files[0], "source_rows": files[1],
                      "valid_rows": files[2], "unique_contacts": sum(counts.values()),
                      "by_status": dict(counts)}, indent=2))


def remap_unresolved(connection, resolver, normalize):
    changed = 0
    rows = connection.execute(
        "SELECT id, record_json FROM contacts WHERE merged_into IS NULL "
        "AND specialty_id IS NULL"
    ).fetchall()
    with connection:
        for item in rows:
            record = json.loads(item["record_json"])
            mapping = resolver.resolve(record["sourceProfession"],
                                       record["sourceSpecialty"],
                                       record["originalSpecialty"])
            if not mapping:
                continue
            record["nexusMapping"] = mapping
            connection.execute(
                "UPDATE contacts SET record_json=?, score=?, profession_id=?, "
                "specialty_id=? WHERE id=?",
                (json.dumps(record, ensure_ascii=False), record_score(record),
                 mapping["professionId"], mapping["specialtyId"], item["id"]),
            )
            changed += 1
    print(f"newly resolved: {changed}", flush=True)
    show_status(connection)


def matching_candidate(client, record, normalize, lookup_client=None):
    matches = []
    matched_client = None
    for search_client in (lookup_client, client):
        if search_client is None:
            continue
        for field in ("email", "phone"):
            if record[field]:
                matches.extend(search_client.search_candidates(**{field: record[field]}))
        if matches:
            matched_client = search_client
            break
    if not matches:
        return None, "", None
    named = [item for item in matches
             if normalize(item.get("firstName"))[:1] == normalize(record["firstName"])[:1]
             and normalize(item.get("lastName")) == normalize(record["lastName"])]
    ids = {item.get("candidateId", item.get("id")) for item in named}
    if len(ids) != 1:
        return None, "existing Nexus contact needs identity review", None
    return named[0], "", matched_client


def process_contact(record, defaults, client, normalize, csv_import, unknown_ids,
                    lookup_client=None):
    match, reason, matched_client = matching_candidate(
        client, record, normalize, lookup_client)
    if reason:
        return "review", None, reason
    mapping = record["nexusMapping"]
    if match:
        candidate_id = match.get("candidateId", match.get("id"))
        existing_ids = {int(value) for value in (match.get("specialtyIds") or [])}
        primary_id = match.get("primarySpecialtyId")
        if primary_id is not None:
            existing_ids.add(int(primary_id))
        if mapping["specialtyId"] in existing_ids:
            return "already_correct", candidate_id, ""
        if existing_ids and not existing_ids <= unknown_ids:
            return "review", candidate_id, "existing non-Unknown specialty differs"
        response = matched_client.patch_candidate(candidate_id, {
            "professionIds": [mapping["professionId"]],
            "specialtyIds": [mapping["specialtyId"]],
            "primarySpecialtyId": mapping["specialtyId"],
        })
        if response.status_code >= 400:
            return "failed", candidate_id, f"Nexus PATCH returned {response.status_code}"
        return "updated", candidate_id, ""
    payload = csv_import.candidate_payload(record, defaults)
    try:
        response = client.create_candidate(payload)
    except Exception as exc:
        if getattr(exc, "status_code", None) == 400:
            try:
                error_code = json.loads(exc.body or "{}").get("errorCode")
            except ValueError:
                error_code = None
            if error_code == "duplicate.candidate.found":
                return "review", None, "Nexus reports a duplicate not returned by candidate search"
        raise
    try:
        body = response.json()
    except ValueError:
        body = {}
    candidate_id = body.get("id", body.get("Id")) if isinstance(body, dict) else None
    if candidate_id is None:
        return "failed", None, "Nexus returned no candidate ID after creation"
    return "created", candidate_id, ""


def run(connection, main, client, nexus_error, csv_import, normalize,
        limit, retry_failed, lookup_client=None):
    rows = connection.execute("SELECT DISTINCT state_code FROM contacts "
                              "WHERE merged_into IS NULL AND conflict=0 "
                              "AND specialty_id IS NOT NULL").fetchall()
    if not rows:
        print("No mapped contacts are ready")
        return
    job_type = main.config.NEXUS_DEFAULT_PROFILE.get("jobTypeIds", ["TRAVEL"])[0]
    defaults = main._csv_candidate_defaults(
        job_type, {row[0] for row in rows}, None)
    unknown_ids = {int(row.get("specialtyId", row.get("id")))
                   for row in client.get_master_cached("specialties")
                   if normalize(row.get("name")) == "unknown"}
    processed = 0
    last_id = 0
    allowed_statuses = ("pending", "failed") if retry_failed else ("pending",)
    while limit is None or processed < limit:
        batch = connection.execute(
            "SELECT id, record_json FROM contacts WHERE id>? AND merged_into IS NULL "
            "AND conflict=0 AND specialty_id IS NOT NULL AND status IN ("
            + ",".join("?" for _ in allowed_statuses) + ") ORDER BY id LIMIT ?",
            (last_id, *allowed_statuses,
             min(100, (limit - processed) if limit else 100))
        ).fetchall()
        if not batch:
            break
        for item in batch:
            last_id = item["id"]
            record = json.loads(item["record_json"])
            for attempt in range(3):
                try:
                    status, candidate_id, error = process_contact(
                        record, defaults, client, normalize, csv_import,
                        unknown_ids, lookup_client)
                    break
                except nexus_error as exc:
                    if attempt == 2 or exc.status_code not in (None, 429, 500, 502, 503, 504):
                        status, candidate_id, error = "failed", None, str(exc)[:300]
                        break
                    time.sleep(2 ** attempt)
                except Exception as exc:
                    status, candidate_id, error = "failed", None, str(exc)[:300]
                    break
            with connection:
                connection.execute(
                    "UPDATE contacts SET status=?, candidate_id=?, error=? WHERE id=?",
                    (status, candidate_id, error, item["id"]))
            processed += 1
            if processed % 100 == 0:
                print(f"processed {processed} in this run", flush=True)
    show_status(connection)


def main_cli():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("prepare", "remap", "run", "status"))
    parser.add_argument("--downloads", type=Path, default=Path.home() / "Downloads")
    parser.add_argument("--env", type=Path, default=Path.home() / "Downloads" / "Nexus_app.env")
    parser.add_argument("--lookup-env", type=Path,
                        default=Path.home() / "Downloads" / "HealthcareBoard.env")
    parser.add_argument("--database", type=Path,
                        default=Path.home() / "AppData" / "Local" / "NexusUploader"
                        / "specialty_import.sqlite3")
    parser.add_argument("--limit", type=int)
    parser.add_argument("--retry-failed", action="store_true")
    args = parser.parse_args()
    connection = open_database(args.database)
    if args.command == "status":
        show_status(connection)
        return
    csv_import, client, nexus_error, resolver_class, normalize = load_app(args.env)
    if args.command == "prepare":
        resolver = resolver_class(client.get_master_cached("professions"),
                                  client.get_master_cached("specialties"))
        prepare(connection, args.downloads, resolver, csv_import, normalize)
        show_status(connection)
    elif args.command == "remap":
        resolver = resolver_class(client.get_master_cached("professions"),
                                  client.get_master_cached("specialties"))
        remap_unresolved(connection, resolver, normalize)
    else:
        if connection.execute("SELECT COUNT(*) FROM files").fetchone()[0] != len(WORKBOOK_NAMES):
            raise RuntimeError("Prepare all requested workbooks before uploading")
        from app import main
        logging.getLogger("httpx").setLevel(logging.WARNING)
        lookup_client = load_lookup_client(args.lookup_env)
        run(connection, main, client, nexus_error, csv_import, normalize,
            args.limit, args.retry_failed, lookup_client)


if __name__ == "__main__":
    main_cli()
